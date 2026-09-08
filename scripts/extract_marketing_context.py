from pathlib import Path
import json
import pandas as pd
from docx import Document
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "marketing_10_day_plan" / "extracted_context.json"

result = {"spreadsheets": {}, "docx": {}, "pptx": {}}

for name in [
    "Thực Phẩm Số 1.xlsx",
    "DanhSachSanPham_KV23062026-132405-026.xlsx",
    "CustomerProfit (21).xls",
    "dulieu/Thực Phẩm Số Một++2_Chiến dịch+3_Nhóm quảng cáo+2026-07-08.csv",
]:
    path = ROOT / name
    try:
        if path.suffix.lower() == ".csv":
            encoding = "utf-16" if path.read_bytes()[:2] in (b"\xff\xfe", b"\xfe\xff") else "utf-8-sig"
            df = pd.read_csv(path, encoding=encoding, sep=None, engine="python")
            result["spreadsheets"][name] = {"shape": list(df.shape), "columns": list(map(str, df.columns)), "sample": df.head(12).fillna("").astype(str).to_dict("records")}
        else:
            book = pd.ExcelFile(path)
            sheets = {}
            for sheet in book.sheet_names:
                df = pd.read_excel(path, sheet_name=sheet)
                sheets[sheet] = {"shape": list(df.shape), "columns": list(map(str, df.columns)), "sample": df.head(8).fillna("").astype(str).to_dict("records")}
            result["spreadsheets"][name] = sheets
    except Exception as exc:
        result["spreadsheets"][name] = {"error": str(exc)}

for name in ["LUỒNG HIỆN TẠI.docx"]:
    try:
        doc = Document(ROOT / name)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        tables = []
        for table in doc.tables:
            tables.append([[cell.text.strip() for cell in row.cells] for row in table.rows])
        result["docx"][name] = {"paragraphs": paragraphs, "tables": tables}
    except Exception as exc:
        result["docx"][name] = {"error": str(exc)}

for name in ["Profile TPS1ĐN-2025 .pptx"]:
    try:
        prs = Presentation(ROOT / name)
        slides = []
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
            slides.append({"slide": i, "text": texts})
        result["pptx"][name] = slides
    except Exception as exc:
        result["pptx"][name] = {"error": str(exc)}

OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(OUT)
